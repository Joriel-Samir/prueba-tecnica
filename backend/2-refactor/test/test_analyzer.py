import os
from unittest.mock import patch as mock_patch

import openpyxl
import pytest

from src.analyzer import ExcelSentimentAnalyzer
from src.main import main
from src.provider import (
    ParallelDotsProvider,
    SentimentAnalysisError,
    SentimentProvider,
)

# Usamos una API key falsa para simular el comportamiento del proveedor sin hacer llamadas reales a la API.
FAKE_API_KEY = "unit-test-sin-acceso-real"


class MockSentimentProvider(SentimentProvider):
    """Doble de prueba para aislar el test de llamadas reales HTTP."""

    def analyze(self, text: str) -> dict[str, float]:
        if "excelente" in text.lower():
            return {"negative": 5.0, "neutral": 10.0, "positive": 85.0}
        if "terrible" in text.lower():
            return {"negative": 90.0, "neutral": 5.0, "positive": 5.0}
        return {"negative": 0.0, "neutral": 0.0, "positive": 0.0}


class FailingSentimentProvider(SentimentProvider):
    """Doble que siempre falla, para probar que no se escriben ceros falsos."""

    def analyze(self, text: str) -> dict[str, float]:
        raise SentimentAnalysisError("fallo simulado")


@pytest.fixture
def dummy_excel(tmp_path):
    file_path = tmp_path / "input_test.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.cell(1, 3).value = "Comentario"
    sheet.cell(2, 3).value = "El producto es excelente"
    sheet.cell(3, 3).value = "La atención fue terrible"
    sheet.cell(4, 3).value = None  # Fila vacía
    workbook.save(file_path)
    return str(file_path)


def test_excel_analyzer_success(dummy_excel, tmp_path):
    output_path = str(tmp_path / "output_test.xlsx")
    provider = MockSentimentProvider()
    analyzer = ExcelSentimentAnalyzer(provider=provider, text_col=3)

    analyzer.process(input_path=dummy_excel, output_path=output_path)

    assert os.path.exists(output_path)

    wb = openpyxl.load_workbook(output_path)
    sheet = wb.active

    # Encabezados
    assert sheet.cell(1, 4).value == "NEGATIVO"
    assert sheet.cell(1, 5).value == "NEUTRAL"
    assert sheet.cell(1, 6).value == "POSITIVO"

    # Fila 2 (Excelente -> Positivo)
    assert sheet.cell(2, 6).value == 85.0

    # Fila 3 (Terrible -> Negativo)
    assert sheet.cell(3, 4).value == 90.0

    # Fila 4 estaba vacía: no se llama al provider, no se escribe nada
    assert sheet.cell(4, 4).value is None
    assert sheet.cell(4, 5).value is None
    assert sheet.cell(4, 6).value is None


def test_excel_analyzer_skips_failed_row_instead_of_writing_zeros(
    dummy_excel, tmp_path
):
    """Si el proveedor falla, la fila queda en blanco (no con 0,0,0 que parece un dato real)."""
    output_path = str(tmp_path / "output_test.xlsx")
    analyzer = ExcelSentimentAnalyzer(provider=FailingSentimentProvider(), text_col=3)

    analyzer.process(input_path=dummy_excel, output_path=output_path)

    sheet = openpyxl.load_workbook(output_path).active
    assert sheet.cell(2, 4).value is None
    assert sheet.cell(2, 5).value is None
    assert sheet.cell(2, 6).value is None


def test_excel_analyzer_file_not_found(tmp_path):
    provider = MockSentimentProvider()
    analyzer = ExcelSentimentAnalyzer(provider=provider)

    with pytest.raises(FileNotFoundError):
        analyzer.process("archivo_inexistente.xlsx", str(tmp_path / "out.xlsx"))


def test_main_loads_api_key_from_dotenv(monkeypatch):
    monkeypatch.delenv("PARALLELDOTS_API_KEY", raising=False)

    with (
        mock_patch("src.main.load_dotenv") as mock_load_dotenv,
        mock_patch("src.main.ParallelDotsProvider") as mock_provider,
        mock_patch("src.main.ExcelSentimentAnalyzer") as mock_analyzer,
        mock_patch("sys.argv", ["main", "input.xlsx"]),
    ):
        mock_load_dotenv.side_effect = lambda: os.environ.__setitem__(
            "PARALLELDOTS_API_KEY", FAKE_API_KEY
        )

        main()

    mock_load_dotenv.assert_called_once_with()
    mock_provider.assert_called_once_with(api_key=FAKE_API_KEY)
    mock_analyzer.return_value.process.assert_called_once_with(
        input_path="input.xlsx",
        output_path="sentimentAnalysis.xlsx",
        sheet_name=None,
    )


def test_parallel_dots_provider_success():
    """El proveedor real convierte las fracciones (0-1) que da la API en porcentajes."""
    mock_response = {"sentiment": {"negative": 0.05, "neutral": 0.10, "positive": 0.85}}

    with mock_patch(
        "src.provider.paralleldots.sentiment", return_value=mock_response
    ) as mock_sentiment:
        provider = ParallelDotsProvider(api_key=FAKE_API_KEY)
        result = provider.analyze("Excelente servicio")

        mock_sentiment.assert_called_once_with("Excelente servicio")

        assert result["positive"] == 85.0
        assert result["neutral"] == 10.0
        assert result["negative"] == 5.0


def test_parallel_dots_provider_failure_and_retry():
    """Si la API falla siempre, se reintenta 3 veces y se lanza un error (nunca ceros inventados)."""
    with (
        mock_patch(
            "src.provider.paralleldots.sentiment", side_effect=Exception("API Down")
        ) as mock_sentiment,
        mock_patch("src.provider.time.sleep"),  # no esperar de verdad en el test
    ):
        provider = ParallelDotsProvider(api_key=FAKE_API_KEY)

        with pytest.raises(SentimentAnalysisError):
            provider.analyze("Texto de prueba")

        assert mock_sentiment.call_count == 3
