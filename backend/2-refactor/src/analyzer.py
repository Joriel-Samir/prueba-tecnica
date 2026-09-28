import logging
import os
from typing import Optional, Tuple
import openpyxl
from src.provider import SentimentProvider

logger = logging.getLogger(__name__)


class ExcelSentimentAnalyzer:
    """Procesador de archivos Excel para análisis de sentimientos."""

    def __init__(
        self,
        provider: SentimentProvider,
        text_col: int = 3,
        out_cols: Tuple[int, int, int] = (4, 5, 6),
    ) -> None:
        self.provider = provider
        self.text_col = text_col
        self.out_cols = out_cols

    def process(
        self,
        input_path: str,
        output_path: str,
        sheet_name: Optional[str] = None,
    ) -> None:
        """Procesa el archivo Excel de entrada y guarda el resultado en output_path."""
        if not os.path.isfile(input_path):
            logger.error("El archivo no existe: %s", input_path)
            raise FileNotFoundError(f"Archivo no encontrado: {input_path}")

        logger.info("Cargando archivo Excel: %s", input_path)
        workbook = openpyxl.load_workbook(input_path)

        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        neg_col, neu_col, pos_col = self.out_cols
        sheet.cell(1, neg_col).value = "NEGATIVO"
        sheet.cell(1, neu_col).value = "NEUTRAL"
        sheet.cell(1, pos_col).value = "POSITIVO"

        max_row = sheet.max_row
        logger.info("Procesando filas de la 2 a la %d...", max_row)

        for row in range(2, max_row + 1):
            cell_value = sheet.cell(row, self.text_col).value
            if not cell_value:
                logger.debug("Fila %d vacía en la columna %d. Omitiendo.", row, self.text_col)
                continue

            text = str(cell_value).strip()
            logger.debug("Procesando fila %d: '%s'", row, text[:30])

            sentiments = self.provider.analyze(text)

            sheet.cell(row, neg_col).value = sentiments["negative"]
            sheet.cell(row, neu_col).value = sentiments["neutral"]
            sheet.cell(row, pos_col).value = sentiments["positive"]

        workbook.save(output_path)
        logger.info("Proceso completado con éxito. Archivo guardado en: %s", output_path)